#!/usr/bin/env python3
"""Gera apresentação corporativa MedicFlow-AI (PPTX + PDF + ZIP)."""

from __future__ import annotations

import os
import shutil
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SCREENSHOTS = DOCS / "screenshots"
OUTPUT_PPTX = DOCS / "MEDICFLOW_PRESENTACAO_CORPORATIVA.pptx"
OUTPUT_PDF = DOCS / "MEDICFLOW_PRESENTACAO_CORPORATIVA.pdf"
OUTPUT_ZIP = DOCS / "MEDICFLOW_APRESENTACAO_PACKAGE.zip"
LOGO_PATH = DOCS / "medicflow-logo.png"

# Paleta institucional
BLUE_DARK = RGBColor(0x1E, 0x3A, 0x5F)
BLUE_MID = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x0D, 0x94, 0x88)
GREEN_LIGHT = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x64, 0x74, 0x8B)
GRAY_DARK = RGBColor(0x33, 0x41, 0x55)
GRAY_LIGHT = RGBColor(0xF1, 0xF5, 0xF9)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def extract_logo() -> Path:
    login = SCREENSHOTS / "01-login.png"
    if not login.exists():
        login = DOCS / "evidence" / "01-login.png"
    img = Image.open(login)
    w, h = img.size
    # Logo no painel direito do login (aprox.)
    logo = img.crop((int(w * 0.52), int(h * 0.06), int(w * 0.95), int(h * 0.28)))
    logo.save(LOGO_PATH)
    return LOGO_PATH


def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def add_rect(slide, left, top, width, height, fill_rgb, line_rgb=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.color.rgb = line_rgb
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=GRAY_DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, items, size=14, color=GRAY_DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return box


def add_header_bar(slide, title: str, subtitle: str = ""):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.15), BLUE_DARK)
    add_rect(slide, Inches(0), Inches(1.15), SLIDE_W, Inches(0.06), GREEN)
    add_textbox(slide, Inches(0.6), Inches(0.22), Inches(10), Inches(0.55), title, size=28, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.6), Inches(0.72), Inches(10), Inches(0.35), subtitle, size=13, color=GREEN_LIGHT)
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(11.2), Inches(0.15), height=Inches(0.85))


def add_footer(slide, text: str = "MedicFlow-AI · Plataforma Operacional · 08/06/2026"):
    add_rect(slide, Inches(0), Inches(7.1), SLIDE_W, Inches(0.4), GRAY_LIGHT)
    add_textbox(slide, Inches(0.6), Inches(7.15), Inches(12), Inches(0.3), text, size=9, color=GRAY)


def add_screenshot(slide, filename: str, left, top, width, height=None):
    path = SCREENSHOTS / filename
    if not path.exists():
        path = DOCS / "evidence" / filename
    if path.exists():
        if height:
            slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        else:
            slide.shapes.add_picture(str(path), left, top, width=width)
        return True
    add_textbox(slide, left, top, width, Inches(1), f"[Screenshot: {filename}]", size=12, color=GRAY)
    return False


def add_metric_cards(slide, metrics: list[tuple[str, str]], left, top):
    card_w = Inches(2.15)
    card_h = Inches(1.35)
    gap = Inches(0.2)
    for i, (value, label) in enumerate(metrics):
        x = left + i * (card_w + gap)
        add_rect(slide, x, top, card_w, card_h, WHITE, BLUE_MID)
        add_textbox(slide, x + Inches(0.15), top + Inches(0.2), card_w - Inches(0.3), Inches(0.6), value, size=32, bold=True, color=BLUE_DARK, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.1), top + Inches(0.82), card_w - Inches(0.2), Inches(0.45), label, size=11, color=GRAY, align=PP_ALIGN.CENTER)


def slide_capa(prs: Presentation):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), Inches(5.2), SLIDE_H, BLUE_DARK)
    add_rect(slide, Inches(5.2), Inches(0), Inches(8.13), SLIDE_H, WHITE)
    # Gradiente simulado
    for i in range(8):
        g = int(0x1E + i * 4)
        add_rect(slide, Inches(0), Inches(5.5 + i * 0.25), Inches(5.2), Inches(0.25), RGBColor(0x0D, g, 0x88 - i * 2))

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(7.0), Inches(1.0), width=Inches(3.8))

    add_textbox(slide, Inches(0.55), Inches(0.5), Inches(4.5), Inches(0.4), "MedicFlow-AI · Plataforma Operacional", size=12, color=GREEN_LIGHT)
    add_textbox(slide, Inches(0.55), Inches(2.0), Inches(4.5), Inches(1.8), "Inteligência que conecta.\nOperação que transforma.", size=26, bold=True, color=WHITE)
    add_textbox(slide, Inches(0.55), Inches(4.2), Inches(4.5), Inches(1.0), "Escalas, plantões, faturamento TISS e indicadores em tempo real — em um único fluxo.", size=13, color=WHITE)

    add_textbox(slide, Inches(6.5), Inches(4.5), Inches(6), Inches(0.8), "Apresentação Corporativa", size=36, bold=True, color=BLUE_DARK, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(6.5), Inches(5.3), Inches(6), Inches(0.5), "Visão executiva para diretoria, investidores e sponsors de implantação", size=14, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(6.5), Inches(6.2), Inches(6), Inches(0.4), "Junho 2026 · Base: auditoria funcional do produto", size=11, color=GRAY, align=PP_ALIGN.CENTER)


def slide_quem_somos(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Quem Somos", "Plataforma operacional hospitalar multi-tenant")
    items = [
        "O MedicFlow-AI unifica gestão de escalas e plantões, central operacional com indicadores em tempo real, faturamento TISS (MVP), fechamento financeiro e implantação assistida.",
        "Stack: React 19 + TanStack Start (SSR) + Supabase (Auth, Postgres, Realtime, Storage) + deploy Cloudflare Workers.",
        "Posicionamento V1: plataforma operacional + faturamento TISS MVP — não é prontuário eletrônico nem ERP hospitalar completo.",
        "Validado em staging (medicflow.app.br) com 20 rotas, 17 telas autenticadas e 52 funcionalidades implementadas.",
    ]
    add_bullets(slide, Inches(0.7), Inches(1.6), Inches(5.8), Inches(5.2), items, size=15)
    add_screenshot(slide, "02-dashboard.png", Inches(6.8), Inches(1.55), Inches(5.9), Inches(5.0))
    add_footer(slide)


def slide_problema(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Problema do Mercado", "Desafios operacionais e financeiros em instituições de saúde")
    problems = [
        ("Escalas em planilhas", "Confirmações por WhatsApp, sem visibilidade de cobertura em tempo real."),
        ("Faturamento disperso", "Guias TISS manuais, lotes sem rastreio, glosas sem controle centralizado."),
        ("Fechamento fragmentado", "Competências sem snapshot auditável, conciliação em múltiplas ferramentas."),
        ("Implantação imprevisível", "Demos desorganizadas, sem roteiro, sem validação técnica de go-live."),
        ("Sistemas isolados", "Operação, financeiro e faturamento sem fluxo integrado multi-instituição."),
    ]
    y = Inches(1.55)
    for title, desc in problems:
        add_rect(slide, Inches(0.7), y, Inches(0.12), Inches(0.55), GREEN)
        add_textbox(slide, Inches(1.0), y, Inches(11), Inches(0.35), title, size=16, bold=True, color=BLUE_DARK)
        add_textbox(slide, Inches(1.0), y + Inches(0.38), Inches(11), Inches(0.4), desc, size=13, color=GRAY)
        y += Inches(1.05)
    add_footer(slide)


def slide_solucao(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Solução MedicFlow-AI", "Um único fluxo para operação, financeiro e implantação")
    pillars = [
        "Gestão de escalas (14 dias) e plantões com aceite, recusa e swaps",
        "Central operacional com alertas em tempo real e ações contextuais",
        "Ciclo TISS integrado: convênios → guias → lotes → glosas → repasses",
        "Fechamento financeiro com snapshot, trava e dashboard executivo",
        "Implantação assistida: piloto, demo guiada (7 passos), smoke tests",
    ]
    add_bullets(slide, Inches(0.7), Inches(1.55), Inches(5.5), Inches(4.5), pillars, size=14)
    add_screenshot(slide, "01-login.png", Inches(6.5), Inches(1.55), Inches(6.0), Inches(5.0))
    add_footer(slide)


def slide_publico(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Público-Alvo", "5 perfis RBAC com capabilities granulares")
    rows = [
        ("Hospitais e clínicas", "tenant_admin", "Branding, parametrização, go-live"),
        ("Coordenadores de escala", "coordinator", "Publicar turnos, resolver conflitos, swaps"),
        ("Profissionais de plantão", "professional", "Confirmar plantões, disponibilidade, trocas"),
        ("Equipe financeira", "financial", "TISS, repasses, fechamento, conciliação"),
        ("Diretoria", "tenant_admin + executivo", "KPIs financeiros e operacionais"),
    ]
    y = Inches(1.7)
    add_rect(slide, Inches(0.7), y, Inches(11.8), Inches(0.45), BLUE_DARK)
    for col, x in [("Segmento", 0.8), ("Perfil RBAC", 4.5), ("Necessidade", 7.2)]:
        add_textbox(slide, Inches(x), y + Inches(0.05), Inches(3), Inches(0.35), col, size=12, bold=True, color=WHITE)
    y += Inches(0.55)
    for i, (seg, perfil, nec) in enumerate(rows):
        bg = GRAY_LIGHT if i % 2 == 0 else WHITE
        add_rect(slide, Inches(0.7), y, Inches(11.8), Inches(0.7), bg)
        add_textbox(slide, Inches(0.8), y + Inches(0.12), Inches(3.5), Inches(0.5), seg, size=12, bold=True, color=BLUE_DARK)
        add_textbox(slide, Inches(4.5), y + Inches(0.12), Inches(2.5), Inches(0.5), perfil, size=11, color=GREEN)
        add_textbox(slide, Inches(7.2), y + Inches(0.12), Inches(5), Inches(0.5), nec, size=12, color=GRAY_DARK)
        y += Inches(0.72)
    add_textbox(slide, Inches(0.7), Inches(6.0), Inches(11.5), Inches(0.8), "V1 não atende: prontuário eletrônico, gestão de pacientes, agenda ambulatorial ou envio automático TISS a operadoras.", size=11, color=GRAY)
    add_footer(slide)


def slide_arquitetura(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Arquitetura da Plataforma", "Stack moderna, multi-tenant nativo, deploy flexível")
    add_metric_cards(slide, [
        ("20", "Rotas"),
        ("17", "Telas"),
        ("12", "Fluxos"),
        ("5", "Perfis RBAC"),
        ("52", "Funcionalidades"),
    ], Inches(0.7), Inches(1.55))

    layers = [
        ("Frontend", "React 19 + TanStack Start (SSR) · Roteamento file-based"),
        ("Backend", "~100+ Server Functions (TanStack createServerFn)"),
        ("Dados", "Supabase Postgres + RLS · 30 migrations SQL"),
        ("Realtime", "Supabase Realtime · Central operacional ao vivo"),
        ("Deploy", "Cloudflare Workers (primário) ou Vercel"),
        ("Segurança", "Auth multi-tenant · RBAC · anti brute-force · auditoria"),
    ]
    y = Inches(3.2)
    for name, desc in layers:
        add_rect(slide, Inches(0.7), y, Inches(1.6), Inches(0.5), BLUE_MID)
        add_textbox(slide, Inches(0.75), y + Inches(0.08), Inches(1.5), Inches(0.35), name, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(2.5), y + Inches(0.08), Inches(9.8), Inches(0.35), desc, size=12, color=GRAY_DARK)
        y += Inches(0.62)
    add_footer(slide)


def slide_feature(prs: Presentation, title: str, subtitle: str, bullets: list[str], screenshot: str):
    slide = blank_slide(prs)
    add_header_bar(slide, title, subtitle)
    add_bullets(slide, Inches(0.7), Inches(1.55), Inches(5.2), Inches(5.0), bullets, size=13)
    add_screenshot(slide, screenshot, Inches(6.3), Inches(1.5), Inches(6.3), Inches(5.1))
    add_footer(slide)


def slide_seguranca(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Segurança", "Proteção em todas as camadas — rota, serviço e banco")
    items = [
        "Login multi-tenant com seleção de instituição e recuperação de senha",
        "RBAC com 5 papéis e 30+ capabilities — enforcement em rota + serviço + RLS",
        "Row Level Security (RLS) Postgres em todas as tabelas multi-tenant",
        "Gate anti brute-force e auditoria de login (security audit writer)",
        "Guard de rotas e sessão com Supabase Auth",
        "Isolamento por tenant: branding, dados e operações segregados",
    ]
    add_bullets(slide, Inches(0.7), Inches(1.6), Inches(5.5), Inches(4.5), items, size=14)

    badges = [
        ("RLS", "Postgres"),
        ("RBAC", "5 perfis"),
        ("Auth", "Multi-tenant"),
        ("Audit", "Login trail"),
    ]
    x = Inches(7.0)
    for label, sub in badges:
        add_rect(slide, x, Inches(2.0), Inches(2.5), Inches(1.2), BLUE_DARK)
        add_textbox(slide, x + Inches(0.2), Inches(2.25), Inches(2.1), Inches(0.5), label, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.2), Inches(2.75), Inches(2.1), Inches(0.35), sub, size=11, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)
        x += Inches(2.8)

    add_textbox(slide, Inches(7.0), Inches(5.5), Inches(5.5), Inches(1.0), "Pendente V1: error tracking externo (Sentry/Datadog) — atualmente stub console.", size=11, color=GRAY)
    add_footer(slide)


def slide_implantacao(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Implantação", "Piloto assistido com demo guiada e go-live controlado")
    phases = [
        ("Dia 0", "Branding e parametrização", "/instituicao"),
        ("Dia 0", "Seed demo (ambiente controlado)", "/instituicao"),
        ("Dia 1–2", "Cadastro de convênios TISS", "/tiss"),
        ("Dia 1–2", "Provisionamento RBAC", "Supabase Auth"),
        ("Dia 3", "Abertura de competência financeira", "/financeiro/fechamento"),
        ("Go-live", "Smoke tests + monitoramento", "/lancamento → /operacao"),
    ]
    y = Inches(1.6)
    for fase, acao, rota in phases:
        add_rect(slide, Inches(0.7), y, Inches(1.1), Inches(0.55), GREEN)
        add_textbox(slide, Inches(0.75), y + Inches(0.1), Inches(1.0), Inches(0.35), fase, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(2.0), y + Inches(0.08), Inches(5.5), Inches(0.4), acao, size=13, bold=True, color=BLUE_DARK)
        add_textbox(slide, Inches(7.8), y + Inches(0.08), Inches(4.5), Inches(0.4), rota, size=11, color=GRAY)
        y += Inches(0.72)

    add_textbox(slide, Inches(0.7), Inches(6.0), Inches(5.5), Inches(0.8), "Demo guiada: 7 passos — piloto → home → central → executivo → TISS → instituição → operação", size=11, color=GRAY_DARK)
    add_screenshot(slide, "08-administracao-piloto.png", Inches(7.5), Inches(1.55), Inches(5.0), Inches(4.8))
    add_footer(slide)


def slide_diferenciais(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Diferenciais", "Vantagens competitivas comprovadas no produto")
    diffs = [
        ("Multi-tenant nativo", "RLS Postgres + branding white-label por instituição"),
        ("Operação em tempo real", "Supabase Realtime + Central Operacional"),
        ("RBAC granular", "5 papéis, 30+ capabilities, enforcement triplo"),
        ("Ciclo TISS integrado", "Guias → Lotes → Glosas → Repasses no mesmo tenant"),
        ("Implantação assistida", "Piloto, demo 7 passos, smoke tests, checklist go-live"),
        ("IA operacional (opcional)", "Copilot GPT, agentes, orquestração supervisionada"),
        ("Observabilidade", "Health checks, export diagnóstico, timeline de eventos"),
        ("Deploy flexível", "Cloudflare Workers ou Vercel"),
    ]
    col1, col2 = diffs[:4], diffs[4:]
    y = Inches(1.6)
    for title, desc in col1:
        add_rect(slide, Inches(0.7), y, Inches(0.08), Inches(0.7), GREEN)
        add_textbox(slide, Inches(0.95), y, Inches(5.5), Inches(0.3), title, size=13, bold=True, color=BLUE_DARK)
        add_textbox(slide, Inches(0.95), y + Inches(0.32), Inches(5.5), Inches(0.35), desc, size=11, color=GRAY)
        y += Inches(0.95)
    y = Inches(1.6)
    for title, desc in col2:
        add_rect(slide, Inches(6.8), y, Inches(0.08), Inches(0.7), GREEN)
        add_textbox(slide, Inches(7.05), y, Inches(5.5), Inches(0.3), title, size=13, bold=True, color=BLUE_DARK)
        add_textbox(slide, Inches(7.05), y + Inches(0.32), Inches(5.5), Inches(0.35), desc, size=11, color=GRAY)
        y += Inches(0.95)
    add_footer(slide)


def slide_roadmap(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Roadmap", "GAPs identificados na auditoria — não são promessas de entrega")
    sections = [
        ("Curto prazo — completar V1", [
            "XML TISS conforme ANS (parcial → alta prioridade)",
            "Hub financeiro com KPIs reais (parcial → média)",
            "Error tracking Sentry/Datadog (não implementado)",
            "Cadastro UI de profissionais (não implementado)",
        ]),
        ("Médio prazo", [
            "Envio TISS para operadoras",
            "Conciliação bancária OFX/CNAB",
            "Webhooks de retorno de operadoras",
            "Auto-cadastro / convite de usuários",
        ]),
        ("Longo prazo", [
            "Módulo de Pacientes",
            "Prontuário eletrônico",
            "Agenda ambulatorial",
            "Integração ERP/DRE",
        ]),
    ]
    x_positions = [Inches(0.7), Inches(4.6), Inches(8.5)]
    for (title, items), x in zip(sections, x_positions):
        add_rect(slide, x, Inches(1.55), Inches(3.6), Inches(0.45), BLUE_DARK)
        add_textbox(slide, x + Inches(0.1), Inches(1.6), Inches(3.4), Inches(0.35), title, size=11, bold=True, color=WHITE)
        add_bullets(slide, x + Inches(0.1), Inches(2.15), Inches(3.4), Inches(4.5), [f"• {i}" for i in items], size=10)
    add_footer(slide)


def slide_resultados(prs: Presentation):
    slide = blank_slide(prs)
    add_header_bar(slide, "Resultados Esperados", "Benefícios mensuráveis para a instituição")
    results = [
        ("Redução de planilhas", "Escalas e plantões centralizados em plataforma única"),
        ("Visibilidade de cobertura", "Central com alertas em tempo real e ações direcionadas"),
        ("Rastreabilidade financeira", "Fechamento com snapshot, trava e auditoria de competência"),
        ("Ciclo TISS documentado", "Guias, lotes e glosas rastreáveis no mesmo sistema"),
        ("Implantação previsível", "Checklist, smoke tests e demo guiada de 7 passos"),
        ("White-label rápido", "Branding sem rebuild — logo, cores e banner por tenant"),
        ("Segurança multi-tenant", "RLS + RBAC em todas as camadas de acesso"),
    ]
    y = Inches(1.6)
    for title, desc in results:
        add_rect(slide, Inches(0.7), y, Inches(2.8), Inches(0.55), GREEN)
        add_textbox(slide, Inches(0.8), y + Inches(0.1), Inches(2.6), Inches(0.35), title, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(3.7), y + Inches(0.1), Inches(8.5), Inches(0.35), desc, size=13, color=GRAY_DARK)
        y += Inches(0.72)

    add_metric_cards(slide, [
        ("68%", "Implementado"),
        ("18%", "Parcial"),
        ("14%", "Pendente"),
        ("9", "GAPs críticos"),
        ("12", "Screenshots"),
    ], Inches(0.7), Inches(5.8))
    add_footer(slide)


def slide_encerramento(prs: Presentation):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, BLUE_DARK)
    add_rect(slide, Inches(0), Inches(3.5), SLIDE_W, Inches(0.08), GREEN)

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(5.0), Inches(0.8), width=Inches(3.2))

    add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10.3), Inches(1.0), "MedicFlow-AI entrega valor imediato para gestão de plantões, visibilidade operacional e faturamento TISS básico.", size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(2.0), Inches(3.2), Inches(9.3), Inches(0.8), "Implantação assistida (piloto, demo, go-live) é diferencial comercial sólido para diretoria e sponsors.", size=15, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)

    add_metric_cards(slide, [
        ("20", "Rotas"),
        ("17", "Telas"),
        ("12", "Fluxos"),
        ("5", "RBAC"),
        ("52", "Features"),
    ], Inches(1.8), Inches(4.2))

    add_textbox(slide, Inches(2.0), Inches(6.0), Inches(9.3), Inches(0.6), "medicflow.app.br · Documentação: docs/ · Auditoria: 08/06/2026", size=12, color=WHITE, align=PP_ALIGN.CENTER)


def build_presentation() -> Presentation:
    print("Extraindo logo...")
    extract_logo()

    print("Construindo slides...")
    prs = new_prs()
    slide_capa(prs)
    slide_quem_somos(prs)
    slide_problema(prs)
    slide_solucao(prs)
    slide_publico(prs)
    slide_arquitetura(prs)

    slide_feature(prs, "Dashboard", "Visão operacional do dia com KPIs e escala em tempo real", [
        "Saudação personalizada e contadores de plantões abertos",
        "Confirmações pendentes e swaps em destaque",
        "Escala do dia com turnos programados",
        "Links rápidos para Central, Escalas e Plantões",
        "Atualização via Supabase Realtime",
    ], "02-dashboard.png")

    slide_feature(prs, "Gestão de Escalas", "Calendário 14 dias com timeline de turnos", [
        "Barra horizontal dos próximos 14 dias",
        "Filtros: conflitos, abertos, sem confirmação",
        "Criação e edição de escalas (coordenador)",
        "Timeline detalhada por turno e profissional",
        "Integração com alertas da Central Operacional",
    ], "03-agenda-escalas.png")

    slide_feature(prs, "Gestão de Plantões", "Aceitar, recusar e trocar turnos", [
        "Plantões abertos para aceite/recusa",
        "Meus plantões com histórico e status",
        "Swaps: solicitar, aprovar e negar trocas",
        "Fluxo integrado com escalas e central",
        "Notificações operacionais em tempo real",
    ], "04-plantoes.png")

    slide_feature(prs, "Central Operacional", "Command center com alertas e ações contextuais", [
        "KPIs e alertas em tempo real",
        "Ações direcionadas por tipo de alerta",
        "Copilot GPT (opcional, requer API key)",
        "Agentes e orquestração supervisionada",
        "Subscriptions Realtime para cobertura ao vivo",
    ], "11-central-operacional.png")

    slide_feature(prs, "Perfil Profissional", "Disponibilidade e configuração do profissional", [
        "Dados do profissional: nome, e-mail, papel, instituição",
        "Switch de disponibilidade para plantões",
        "Janelas padrão seg–sex 07:00–19:00",
        "Integração com riscos de cobertura na central",
        "Encerramento seguro de sessão",
    ], "10-perfil.png")

    slide_feature(prs, "Financeiro", "Hub financeiro e fechamento operacional", [
        "Hub financeiro com links para sub-módulos",
        "Fechamento de competência com snapshot e trava",
        "Conciliação operacional com importação CSV",
        "Reabertura de competência (admin)",
        "Repasses médicos vinculados à produção TISS",
    ], "05-financeiro.png")

    slide_feature(prs, "Dashboard Executivo", "KPIs financeiros e operacionais reais", [
        "Indicadores consolidados por competência",
        "Visão para diretoria e equipe financeira",
        "Dados reais (não ilustrativos)",
        "Acesso via perfil financial_closing:read",
        "Integrado ao fluxo de fechamento mensal",
    ], "06-relatorios-dashboard-executivo.png")

    slide_feature(prs, "TISS / Faturamento", "Ciclo completo MVP — convênios a repasses", [
        "Convênios, contratos e catálogo TUSS",
        "Guias TISS e lotes com export XML (MVP)",
        "Glosas e recursos (entrada manual)",
        "Produção médica e repasses integrados",
        "Sem envio automático a operadoras (V1)",
    ], "09-tiss.png")

    slide_seguranca(prs)
    slide_implantacao(prs)
    slide_diferenciais(prs)
    slide_roadmap(prs)
    slide_resultados(prs)
    slide_encerramento(prs)

    return prs


def export_pdf(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PPTX para PDF via PowerPoint COM (Windows)."""
    try:
        import comtypes.client
    except ImportError:
        print("comtypes não disponível para export PDF")
        return False

    pptx_abs = str(pptx_path.resolve())
    pdf_abs = str(pdf_path.resolve())

    powerpoint = None
    presentation = None
    try:
        powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
        powerpoint.Visible = 1
        presentation = powerpoint.Presentations.Open(pptx_abs, WithWindow=False)
        presentation.SaveAs(pdf_abs, 32)  # ppSaveAsPDF
        presentation.Close()
        powerpoint.Quit()
        print(f"PDF exportado via PowerPoint: {pdf_path}")
        return pdf_path.exists()
    except Exception as exc:
        print(f"Falha export PDF via PowerPoint: {exc}")
        try:
            if presentation:
                presentation.Close()
            if powerpoint:
                powerpoint.Quit()
        except Exception:
            pass
        return False


def create_zip() -> None:
    doc_files = [
        "APRESENTACAO_EXECUTIVA_MEDICFLOW.md",
        "AUDITORIA_FUNCIONAL_MEDICFLOW.md",
        "MANUAL_PROFISSIONAL_MEDICFLOW.md",
        "MANUAL_INSTITUICAO_MEDICFLOW.md",
        "DOCUMENTACAO_FINAL_STATUS.md",
        "MEDICFLOW_PRESENTACAO_CORPORATIVA.pptx",
        "MEDICFLOW_PRESENTACAO_CORPORATIVA.pdf",
    ]
    if OUTPUT_ZIP.exists():
        OUTPUT_ZIP.unlink()

    with zipfile.ZipFile(OUTPUT_ZIP, "w", zipfile.ZIP_DEFLATED) as zf:
        for doc in doc_files:
            path = DOCS / doc
            if path.exists():
                zf.write(path, f"documentacao/{doc}")

        if LOGO_PATH.exists():
            zf.write(LOGO_PATH, "assets/medicflow-logo.png")

        for png in sorted(SCREENSHOTS.glob("*.png")):
            zf.write(png, f"screenshots/{png.name}")

    print(f"ZIP criado: {OUTPUT_ZIP}")


def format_size(path: Path) -> str:
    if not path.exists():
        return "N/A"
    size = path.stat().st_size
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / (1024 * 1024):.2f} MB"


def main():
    os.chdir(ROOT)
    prs = build_presentation()
    prs.save(str(OUTPUT_PPTX))
    slide_count = len(prs.slides)
    print(f"PPTX salvo: {OUTPUT_PPTX} ({slide_count} slides)")

    pdf_ok = export_pdf(OUTPUT_PPTX, OUTPUT_PDF)
    if not pdf_ok:
        # Fallback: LibreOffice
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            import subprocess
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(DOCS), str(OUTPUT_PPTX)], check=False)
            print("Tentativa via LibreOffice concluída")

    create_zip()

    print("\n=== ENTREGA ===")
    print(f"Slides: {slide_count}")
    print(f"PPTX: {format_size(OUTPUT_PPTX)}")
    print(f"PDF:  {format_size(OUTPUT_PDF)}")
    print(f"ZIP:  {format_size(OUTPUT_ZIP)}")


if __name__ == "__main__":
    main()
